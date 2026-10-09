"""Convert a private Google Drive XLSX export to the versioned migration format.

Receives private XLSX via stdin and writes JSON to a Git-ignored local file.
Never logs spreadsheet values, addresses, customer names, or payment identifiers.
"""
import sys
import json
from pathlib import Path
from datetime import date, datetime, timedelta
from io import BytesIO
from openpyxl import load_workbook

if len(sys.argv) != 3:
    raise RuntimeError("Snapshot ID and title required")
snapshot_id, title = sys.argv[1:]
workbook = load_workbook(BytesIO(sys.stdin.buffer.read()), data_only=True, read_only=True)
keys = {
    'Vendas':'Venda_ID','Itens da Venda':'Item_ID','Recebimentos':'Pagamento_ID',
    'Produção':'Produção_ID','Compras e Despesas':'Mov_ID','Produtos':'Produto_ID',
    'Insumos':'Insumo_ID','Receitas':'Receita_ID','Ajustes de Estoque':'Ajuste_ID',
    'Kits':'Kit_ID','Itens do Kit':'Kit_Item_ID','Cadastros':'Cadastro_ID',
    'Preços de Formato':'Preço_ID','Adicionais':'Adicional_ID',
    'Itens do Adicional':'Adicional_Item_ID','Consumos da Venda':'Consumo_ID'
}
required = {
    'Vendas','Itens da Venda','Recebimentos','Produção','Compras e Despesas','Produtos',
    'Insumos','Kits','Itens do Kit','Adicionais','Itens do Adicional','Cadastros','Preços de Formato'
}
if missing := required.difference(workbook.sheetnames):
    raise RuntimeError("Required source tabs missing: " + ', '.join(sorted(missing)))
date_columns = {'Data','Vencimento','Data da compra','Data do lote informada','Data Venda'}
sheets, views = {}, {}
for page in workbook.worksheets:
    if page.title not in set(keys) | {
        'Movimentações de Capital','Custos de Aquisição','Caixa e Aportes',
        'Acompanhamento Encomendas','Reservas de Embalagem','Inventário 08-10',
        'Controle Simplificado','Ficha de Produção'
    }:
        continue
    rows = page.iter_rows(min_row=4, values_only=True)
    header = [str(item).strip() if item is not None else '' for item in next(rows, ())]
    if not header:
        continue
    results = []
    for cells in rows:
        if not any(value not in ('', None) for value in cells):
            continue
        record = {}
        for index, key in enumerate(header):
            if not key:
                continue
            value = cells[index] if index < len(cells) else None
            if isinstance(value, (date, datetime)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            record[key] = value if value is not None else ''
        if page.title in keys:
            if not str(record.get(keys[page.title], '')).strip():
                continue
        results.append(record)
    if page.title in keys:
        sheets[page.title] = results
    else:
        views[page.title] = results
counts = {name:len(records) for name,records in sheets.items()}
if counts.get('Vendas',0)<25 or counts.get('Recebimentos',0)<24 or counts.get('Produção',0)<45:
    raise RuntimeError("Snapshot misses required reconciled historic records")
cash_sheet = workbook['Caixa e Aportes']
def money_cents(cell):
    value = cash_sheet[cell].value
    if value is None or not isinstance(value,(int,float)):
        raise RuntimeError("Bank snapshot numerical input missing")
    return round(value*100)
raw_date = cash_sheet['B5'].value
if isinstance(raw_date,datetime):
    recorded_date = raw_date.date().isoformat()
elif isinstance(raw_date,date):
    recorded_date = raw_date.isoformat()
elif isinstance(raw_date,(int,float)):
    recorded_date = (date(1899,12,30)+timedelta(days=int(raw_date))).isoformat()
else:
    raise RuntimeError("Bank snapshot date missing")
bank_snapshot = {
    'balanceCents':money_cents('B4'),
    'businessDate':recorded_date,
    'ownerFundedCents':money_cents('B6'),
    'reimbursementDueCents':money_cents('B9'),
    'source':'legacy-bank-audit'
}
out = {
    'schemaVersion':1,'spreadsheetId':snapshot_id,'spreadsheetName':title,
    'timeZone':'America/Sao_Paulo','exportedAt':datetime.now().isoformat(),
    'sheets':sheets,'views':views,'sourceMetadata':{'bankSnapshot':bank_snapshot}
}
destination = Path('tools/migration/legacy-raw.json')
destination.write_text(json.dumps(out,ensure_ascii=False,default=str),encoding='utf8')
print('Private snapshot converted successfully. Domain source counts: '+json.dumps(counts))
